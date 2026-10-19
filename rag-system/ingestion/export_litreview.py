import os
import json
import openpyxl

DOWNLOADS_DIR = r"C:\Users\shwet\Downloads"
OUTPUT_CORPUS = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "corpus", "corpus.jsonl"))

def clean(val):
    if val is None:
        return ""
    text = str(val).strip()
    return "" if text.lower() == "none" else text

def export_machine_learning():
    path = os.path.join(DOWNLOADS_DIR, "Machine leaarning & adoption prediction final.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Lit Review (Transposed)"]
    
    # Map row field names
    fields = {}
    for r in range(1, ws.max_row + 1):
        field_name = clean(ws.cell(r, 1).value)
        if field_name:
            fields[field_name] = r
            
    r_author = fields.get("Author(s) & Year", 6)
    r_title = fields.get("Title", 7)
    r_journal = fields.get("Journal / Source", 8)
    r_obj = fields.get("Research Objective / Aim", 9)
    r_method = fields.get("Methodology / Models Used", 13)
    r_findings = fields.get("Key Findings", 18)
    r_gaps = fields.get("Limitations / Research Gaps", 20)
    r_future = fields.get("Future Research Directions", 21)
    
    chunks = []
    # Columns 2 through 9
    for col in range(2, ws.max_column + 1):
        title = clean(ws.cell(r_title, col).value)
        author = clean(ws.cell(r_author, col).value)
        if not title and not author:
            continue
        journal = clean(ws.cell(r_journal, col).value)
        objective = clean(ws.cell(r_obj, col).value)
        methodology = clean(ws.cell(r_method, col).value)
        findings = clean(ws.cell(r_findings, col).value)
        gaps = clean(ws.cell(r_gaps, col).value)
        future = clean(ws.cell(r_future, col).value)
        
        chunk_lines = [
            f"Title: {title}",
            f"Authors: {author}",
            f"Journal / Source: {journal}",
            f"Objective: {objective}",
            f"Methodology: {methodology}",
            f"Key Findings: {findings}",
            f"Limitations / Gaps: {gaps}",
            f"Future Work: {future}"
        ]
        text = "\n".join(chunk_lines)
        doc_id = f"lit_ml_{col-1:02d}"
        chunks.append({
            "id": doc_id,
            "source": "litreview",
            "section": "ML & Adoption Prediction",
            "title": title or author,
            "text": text
        })
    return chunks

def export_forecasting():
    path = os.path.join(DOWNLOADS_DIR, "Forecasting & Predictive Analytics Final literature matrix.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws_info = wb["1. Paper Info"]
    ws_comp = wb["2. Full Comparison Matrix"]
    ws_gap = wb["3. Gap & Adaptability Summary"]
    
    # Paper Info rows: 2: Title, 3: Authors, 4: Year, 5: Journal / Conference
    # Full Comp rows: 3: Objective, 9: Flowchart, 10: Models, 18: Key Findings, 20: Stated Limitations
    # Gap Summary rows: 11: Adaptability
    chunks = []
    for col in range(2, ws_info.max_column + 1):
        title = clean(ws_info.cell(2, col).value)
        author = clean(ws_info.cell(3, col).value)
        year = clean(ws_info.cell(4, col).value)
        journal = clean(ws_info.cell(5, col).value)
        if not title and not author:
            continue
        full_author = f"{author} ({year})" if year else author
        
        objective = clean(ws_comp.cell(3, col).value)
        method_flow = clean(ws_comp.cell(9, col).value)
        models = clean(ws_comp.cell(10, col).value)
        methodology = f"{models}. Flow: {method_flow}" if method_flow else models
        
        findings = clean(ws_comp.cell(18, col).value)
        gaps = clean(ws_comp.cell(20, col).value)
        future = clean(ws_gap.cell(11, col).value)
        
        chunk_lines = [
            f"Title: {title}",
            f"Authors: {full_author}",
            f"Journal / Source: {journal}",
            f"Objective: {objective}",
            f"Methodology: {methodology}",
            f"Key Findings: {findings}",
            f"Limitations / Gaps: {gaps}",
            f"Future Work / Adaptability: {future}"
        ]
        text = "\n".join(chunk_lines)
        doc_id = f"lit_forecast_{col-1:02d}"
        chunks.append({
            "id": doc_id,
            "source": "litreview",
            "section": "Forecasting & Predictive Analytics",
            "title": title or full_author,
            "text": text
        })
    return chunks

def export_pmsurya_ghar():
    path = os.path.join(DOWNLOADS_DIR, "PM Surya ghar & Indian context final.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws_idx = wb["0. Paper Index"]
    ws_obj = wb["1. Problem & Objectives"]
    ws_meth = wb["5. Data & Methodology"]
    ws_find = wb["8. Key Findings"]
    ws_gap = wb["9. Limitations & Research Gap"]
    
    chunks = []
    # Rows 2 to 6 are P1 to P5
    for r in range(2, ws_idx.max_row + 1):
        p_id = clean(ws_idx.cell(r, 1).value)
        title = clean(ws_idx.cell(r, 2).value)
        author = clean(ws_idx.cell(r, 3).value)
        journal = clean(ws_idx.cell(r, 4).value)
        year = clean(ws_idx.cell(r, 5).value)
        if not title and not author:
            continue
        full_author = f"{author} ({year})" if year else author
        
        problem = clean(ws_obj.cell(r, 2).value)
        obj = clean(ws_obj.cell(r, 3).value)
        objective = f"{obj} (Problem: {problem})" if problem else obj
        
        d_src = clean(ws_meth.cell(r, 2).value)
        m_tech = clean(ws_meth.cell(r, 6).value)
        methodology = f"{m_tech}. Data: {d_src}" if d_src else m_tech
        
        findings = clean(ws_find.cell(r, 2).value)
        gap1 = clean(ws_gap.cell(r, 2).value)
        gap2 = clean(ws_gap.cell(r, 4).value)
        gaps = f"{gap1}. Research Gap: {gap2}" if gap2 else gap1
        future = f"Adaptability to PM Surya Ghar national/state monitoring"
        
        chunk_lines = [
            f"Title: {title}",
            f"Authors: {full_author}",
            f"Journal / Source: {journal}",
            f"Objective: {objective}",
            f"Methodology: {methodology}",
            f"Key Findings: {findings}",
            f"Limitations / Gaps: {gaps}",
            f"Future Work: {future}"
        ]
        text = "\n".join(chunk_lines)
        doc_id = f"lit_pmsg_{r-1:02d}"
        chunks.append({
            "id": doc_id,
            "source": "litreview",
            "section": "PM Surya Ghar & Indian Context",
            "title": title or full_author,
            "text": text
        })
    return chunks

def export_policy_barriers():
    path = os.path.join(DOWNLOADS_DIR, "Policy_Barriers_Socioeconomic_FINAL.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Policy_Barriers_Batch2"]
    
    chunks = []
    # Row 1 is headers, rows 2 to 7 are the 6 papers
    for r in range(2, ws.max_row + 1):
        title = clean(ws.cell(r, 1).value)
        author = clean(ws.cell(r, 2).value)
        year = clean(ws.cell(r, 3).value)
        journal = clean(ws.cell(r, 4).value)
        if not title and not author:
            continue
        full_author = f"{author} ({year})" if year else author
        
        objective = clean(ws.cell(r, 9).value)
        method_flow = clean(ws.cell(r, 19).value)
        models = clean(ws.cell(r, 20).value)
        methodology = f"{models}. Flow: {method_flow}" if method_flow else models
        
        findings1 = clean(ws.cell(r, 15).value)
        findings2 = clean(ws.cell(r, 24).value)
        findings = f"{findings1}. {findings2}" if findings2 else findings1
        
        gaps1 = clean(ws.cell(r, 22).value)
        gaps2 = clean(ws.cell(r, 25).value)
        gaps3 = clean(ws.cell(r, 28).value)
        gaps = f"Barriers: {gaps1}. Limitations: {gaps2}. Gap: {gaps3}"
        
        future = clean(ws.cell(r, 27).value)
        
        chunk_lines = [
            f"Title: {title}",
            f"Authors: {full_author}",
            f"Journal / Source: {journal}",
            f"Objective: {objective}",
            f"Methodology: {methodology}",
            f"Key Findings: {findings}",
            f"Limitations / Gaps: {gaps}",
            f"Future Work: {future}"
        ]
        text = "\n".join(chunk_lines)
        doc_id = f"lit_policy_{r-1:02d}"
        chunks.append({
            "id": doc_id,
            "source": "litreview",
            "section": "Policy, Barriers & Socioeconomic Factors",
            "title": title or full_author,
            "text": text
        })
    return chunks

def export_solar_policy_landscape():
    path = os.path.join(DOWNLOADS_DIR, "Solar energy and policy landscape final.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["India Solar Policy Lit Review"]
    
    # Map row field names
    fields = {}
    for r in range(1, ws.max_row + 1):
        field_name = clean(ws.cell(r, 1).value)
        if field_name:
            fields[field_name] = r
            
    r_author = fields.get("Author(s) & Year", 6)
    r_title = fields.get("Title", 7)
    r_journal = fields.get("Journal / Source", 8)
    r_obj = fields.get("Research Objective / Aim", 9)
    r_method = fields.get("Methodology / Models Used", 15)
    r_findings = fields.get("Key Findings", 20)
    r_gaps = fields.get("Limitations / Research Gaps", 22)
    r_future = fields.get("Future Research Directions", 23)
    
    chunks = []
    # Columns 2 to 7
    for col in range(2, ws.max_column + 1):
        title = clean(ws.cell(r_title, col).value)
        author = clean(ws.cell(r_author, col).value)
        if not title and not author:
            continue
        journal = clean(ws.cell(r_journal, col).value)
        objective = clean(ws.cell(r_obj, col).value)
        methodology = clean(ws.cell(r_method, col).value)
        findings = clean(ws.cell(r_findings, col).value)
        gaps = clean(ws.cell(r_gaps, col).value)
        future = clean(ws.cell(r_future, col).value)
        
        chunk_lines = [
            f"Title: {title}",
            f"Authors: {author}",
            f"Journal / Source: {journal}",
            f"Objective: {objective}",
            f"Methodology: {methodology}",
            f"Key Findings: {findings}",
            f"Limitations / Gaps: {gaps}",
            f"Future Work: {future}"
        ]
        text = "\n".join(chunk_lines)
        doc_id = f"lit_landscape_{col-1:02d}"
        chunks.append({
            "id": doc_id,
            "source": "litreview",
            "section": "Solar Energy & Policy Landscape",
            "title": title or author,
            "text": text
        })
    return chunks

def main():
    os.makedirs(os.path.dirname(OUTPUT_CORPUS), exist_ok=True)
    all_chunks = []
    
    c_ml = export_machine_learning()
    print(f"Exported {len(c_ml)} chunks from ML & Adoption Prediction")
    all_chunks.extend(c_ml)
    
    c_fc = export_forecasting()
    print(f"Exported {len(c_fc)} chunks from Forecasting & Predictive Analytics")
    all_chunks.extend(c_fc)
    
    c_pm = export_pmsurya_ghar()
    print(f"Exported {len(c_pm)} chunks from PM Surya Ghar & Indian Context")
    all_chunks.extend(c_pm)
    
    c_pb = export_policy_barriers()
    print(f"Exported {len(c_pb)} chunks from Policy, Barriers & Socioeconomic Factors")
    all_chunks.extend(c_pb)
    
    c_sp = export_solar_policy_landscape()
    print(f"Exported {len(c_sp)} chunks from Solar Energy & Policy Landscape")
    all_chunks.extend(c_sp)
    
    print(f"Total literature review chunks: {len(all_chunks)}")
    
    # If corpus.jsonl already exists, load existing non-litreview records or overwrite fresh
    existing_records = []
    if os.path.exists(OUTPUT_CORPUS):
        with open(OUTPUT_CORPUS, "r", encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    item = json.loads(line)
                    if item.get("source") != "litreview":
                        existing_records.append(item)
                        
    combined = all_chunks + existing_records
    with open(OUTPUT_CORPUS, "w", encoding="utf-8") as f:
        for item in combined:
            f.write(json.dumps(item, ensure_ascii=False) + "\n")
            
    print(f"Successfully wrote {len(combined)} records to {OUTPUT_CORPUS}")

if __name__ == "__main__":
    main()
